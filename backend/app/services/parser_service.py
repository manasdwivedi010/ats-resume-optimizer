import io
import re
from typing import Dict, Any, List
import pdfplumber
import pypdf
import docx

class ResumeParserService:
    @staticmethod
    def extract_text_from_pdf(file_bytes: bytes) -> str:
        text = ""
        try:
            with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
                for page in pdf.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text += page_text + "\n"
        except Exception:
            reader = pypdf.PdfReader(io.BytesIO(file_bytes))
            for page in reader.pages:
                extracted = page.extract_text()
                if extracted:
                    text += extracted + "\n"
        return text.strip()

    @staticmethod
    def extract_text_from_docx(file_bytes: bytes) -> str:
        doc = docx.Document(io.BytesIO(file_bytes))
        full_text = []
        for para in doc.paragraphs:
            if para.text.strip():
                full_text.append(para.text.strip())
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    full_text.append(" | ".join(row_text))
        return "\n".join(full_text)

    @staticmethod
    def parse_resume(file_bytes: bytes, filename: str) -> Dict[str, Any]:
        ext = filename.lower().split(".")[-1]
        if ext == "pdf":
            raw_text = ResumeParserService.extract_text_from_pdf(file_bytes)
        elif ext in ("docx", "doc"):
            raw_text = ResumeParserService.extract_text_from_docx(file_bytes)
        elif ext == "txt":
            raw_text = file_bytes.decode("utf-8", errors="ignore")
        else:
            raise ValueError(f"Unsupported file format: .{ext}. Please upload a PDF or DOCX file.")

        if not raw_text.strip():
            raise ValueError("Could not extract any readable text from the file. It may be a scanned image or corrupted.")

        # Detect contact details
        email_match = re.findall(r"[\w\.-]+@[\w\.-]+\.\w+", raw_text)
        phone_match = re.findall(r"(?:\+?\d{1,3}[-.\s]?)?\(?\d{3}\)?[-.\s]?\d{3}[-.\s]?\d{4}", raw_text)
        linkedin_match = re.findall(r"linkedin\.com/in/[\w\-]+", raw_text, re.IGNORECASE)
        github_match = re.findall(r"github\.com/[\w\-]+", raw_text, re.IGNORECASE)

        # Detect candidate name from top lines
        candidate_name = None
        for line in raw_text.splitlines()[:6]:
            l = line.strip()
            if l and 2 <= len(l.split()) <= 4 and len(l) < 45:
                if not any(c in l for c in ["@", "http", "www.", ".com", "/", "\\", "|", "+"]):
                    if re.search(r"^[a-zA-Z\s\.\,\-]+$", l):
                        candidate_name = l
                        break

        # Detect sections
        sections = ResumeParserService._extract_sections(raw_text)

        # Word count & formatting checks
        word_count = len(raw_text.split())
        bullet_matches = re.findall(r"(?:^|\n)\s*[\u2022\u2023\u25E6\u2043\u2219\*\-\–]\s+(.+)", raw_text)

        return {
            "raw_text": raw_text,
            "word_count": word_count,
            "bullet_count": len(bullet_matches),
            "contact": {
                "name": candidate_name,
                "email": email_match[0] if email_match else None,
                "phone": phone_match[0] if phone_match else None,
                "linkedin": linkedin_match[0] if linkedin_match else None,
                "github": github_match[0] if github_match else None
            },
            "sections": sections
        }

    @staticmethod
    def _extract_sections(text: str) -> Dict[str, str]:
        """Categorizes text into standard ATS resume sections."""
        section_patterns = {
            "summary": r"(?:summary|professional summary|profile|about me|objective)",
            "skills": r"(?:skills|technical skills|technologies|core competencies|expertise)",
            "experience": r"(?:experience|work experience|employment history|professional experience)",
            "education": r"(?:education|academic background|qualifications)",
            "projects": r"(?:projects|personal projects|key projects)",
            "certifications": r"(?:certifications|licenses|courses|awards)"
        }

        found_sections: Dict[str, str] = {}
        lines = text.splitlines()
        current_section = "general"
        section_buffers: Dict[str, List[str]] = {k: [] for k in section_patterns}
        section_buffers["general"] = []

        for line in lines:
            line_clean = line.strip().lower()
            matched_sec = None
            if len(line_clean) < 40:
                for sec_name, pattern in section_patterns.items():
                    if re.match(r"^[\s#*\-_]*" + pattern + r"[\s:*\-_]*$", line_clean):
                        matched_sec = sec_name
                        break
            
            if matched_sec:
                current_section = matched_sec
            else:
                if current_section in section_buffers:
                    section_buffers[current_section].append(line)

        for sec, buffer in section_buffers.items():
            if buffer:
                found_sections[sec] = "\n".join(buffer).strip()

        return found_sections

parser_service = ResumeParserService()
